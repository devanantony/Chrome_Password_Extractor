#!/usr/bin/env python3
"""
Chrome Passwords CSV Exporter
Reads a Chrome-exported Passwords.csv locally and exports:
  1) credentials.txt — tab-separated Domain, Username, Password
  2) credentials.xlsx — the same data in three columns

Security: exported files contain passwords in plain text. Store them securely
and delete them when no longer needed.
Dependency: openpyxl (install with: python -m pip install openpyxl)
"""

import csv
import os
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from urllib.parse import urlparse
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError:
    Workbook = None


APP_TITLE = "Chrome Passwords Exporter"


def domain_from_url(value):
    value = (value or "").strip()
    if not value:
        return ""
    candidate = value if "://" in value else "https://" + value
    try:
        host = urlparse(candidate).hostname or ""
        return host.lower()
    except ValueError:
        return value


def find_column(headers, candidates):
    normalized = {re.sub(r"[^a-z0-9]", "", h.lower()): h for h in headers}
    for candidate in candidates:
        key = re.sub(r"[^a-z0-9]", "", candidate.lower())
        if key in normalized:
            return normalized[key]
    return None


def read_chrome_csv(path):
    # utf-8-sig handles the BOM that can appear in Chrome exports.
    with open(path, "r", encoding="utf-8-sig", newline="") as file:
        sample = file.read(8192)
        file.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(file, dialect=dialect)
        if not reader.fieldnames:
            raise ValueError("The CSV file appears to be empty or has no header row.")

        headers = [h.strip() for h in reader.fieldnames if h]
        url_col = find_column(headers, ["url", "origin", "website", "domain"])
        user_col = find_column(headers, ["username", "user name", "login username"])
        pass_col = find_column(headers, ["password", "login password"])

        if not url_col or not user_col or not pass_col:
            raise ValueError(
                "Could not identify the required columns.\n\n"
                "Expected headers similar to: url, username, password.\n"
                "Found: " + ", ".join(headers)
            )

        rows = []
        for row in reader:
            # Preserve password text exactly; do not strip it.
            url = row.get(url_col, "") or ""
            username = row.get(user_col, "") or ""
            password = row.get(pass_col, "") or ""
            if not (url or username or password):
                continue
            rows.append({
                "Domain": domain_from_url(url),
                "Username": username,
                "Password": password,
            })
        return rows


def export_txt(rows, output_path):
    with open(output_path, "w", encoding="utf-8", newline="") as file:
        writer = csv.writer(file, delimiter="\t", lineterminator="\n")
        writer.writerow(["Domain", "Username", "Password"])
        for row in rows:
            writer.writerow([row["Domain"], row["Username"], row["Password"]])


def export_xlsx(rows, output_path):
    if Workbook is None:
        raise RuntimeError(
            "The openpyxl package is not installed.\n\n"
            "Install it with:\npython -m pip install openpyxl"
        )
    wb = Workbook()
    ws = wb.active
    ws.title = "Passwords"
    headers = ["Domain", "Username", "Password"]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append([row["Domain"], row["Username"], row["Password"]])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 36
    ws.row_dimensions[1].height = 24

    # Excel table adds filters and readable row styling.
    if rows:
        table = Table(displayName="PasswordExport", ref=f"A1:C{len(rows) + 1}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    wb.save(output_path)


class PasswordExporterApp:
    BG = "#F4F7FB"
    CARD = "#FFFFFF"
    TEXT = "#172033"
    MUTED = "#64748B"
    BLUE = "#2563EB"
    BORDER = "#DCE4EF"

    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("760x520")
        self.root.minsize(680, 480)
        self.root.configure(bg=self.BG)

        self.csv_path = tk.StringVar()
        self.output_dir = tk.StringVar()
        self.status = tk.StringVar(value="Choose your Chrome Passwords.csv file to begin.")
        self.count_text = tk.StringVar(value="No file loaded")
        self.rows = None

        self._configure_style()
        self._build_ui()

    def _configure_style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("TFrame", background=self.BG)
        style.configure("Card.TFrame", background=self.CARD)
        style.configure("TLabel", background=self.BG, foreground=self.TEXT, font=("Segoe UI", 10))
        style.configure("Title.TLabel", background=self.BG, foreground=self.TEXT, font=("Segoe UI", 21, "bold"))
        style.configure("Sub.TLabel", background=self.BG, foreground=self.MUTED, font=("Segoe UI", 10))
        style.configure("Card.TLabel", background=self.CARD, foreground=self.TEXT, font=("Segoe UI", 10))
        style.configure("CardMuted.TLabel", background=self.CARD, foreground=self.MUTED, font=("Segoe UI", 9))
        style.configure("TEntry", padding=9, fieldbackground="#FFFFFF")
        style.configure("TButton", padding=(12, 8), font=("Segoe UI", 10))
        style.configure("Primary.TButton", padding=(14, 10), background=self.BLUE, foreground="#FFFFFF",
                        font=("Segoe UI", 10, "bold"))
        style.map("Primary.TButton", background=[("active", "#1D4ED8"), ("disabled", "#93C5FD")],
                  foreground=[("disabled", "#FFFFFF")])
        style.configure("TLabelframe", background=self.CARD, bordercolor=self.BORDER)
        style.configure("TLabelframe.Label", background=self.CARD, foreground=self.TEXT,
                        font=("Segoe UI", 10, "bold"))

    def _build_ui(self):
        outer = ttk.Frame(self.root, padding=(28, 24, 28, 20))
        outer.pack(fill="both", expand=True)

        ttk.Label(outer, text="Chrome Passwords Exporter", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            outer,
            text="Convert Chrome's CSV export into a tab-separated text file and an Excel workbook.",
            style="Sub.TLabel",
        ).pack(anchor="w", pady=(5, 20))

        card = ttk.Frame(outer, style="Card.TFrame", padding=20)
        card.pack(fill="x")
        ttk.Label(card, text="1  ·  Source CSV", style="Card.TLabel",
                  font=("Segoe UI", 11, "bold")).pack(anchor="w")
        source_row = ttk.Frame(card, style="Card.TFrame")
        source_row.pack(fill="x", pady=(10, 5))
        self.source_entry = ttk.Entry(source_row, textvariable=self.csv_path)
        self.source_entry.pack(side="left", fill="x", expand=True, ipady=2)
        ttk.Button(source_row, text="Browse…", command=self.choose_csv).pack(side="left", padx=(10, 0))
        ttk.Label(card, textvariable=self.count_text, style="CardMuted.TLabel").pack(anchor="w", pady=(4, 0))

        ttk.Separator(outer).pack(fill="x", pady=17)

        card2 = ttk.Frame(outer, style="Card.TFrame", padding=20)
        card2.pack(fill="x")
        ttk.Label(card2, text="2  ·  Output location", style="Card.TLabel",
                  font=("Segoe UI", 11, "bold")).pack(anchor="w")
        output_row = ttk.Frame(card2, style="Card.TFrame")
        output_row.pack(fill="x", pady=(10, 5))
        ttk.Entry(output_row, textvariable=self.output_dir).pack(side="left", fill="x", expand=True, ipady=2)
        ttk.Button(output_row, text="Choose folder…", command=self.choose_output).pack(side="left", padx=(10, 0))
        ttk.Label(card2, text="Creates credentials.txt and credentials.xlsx in this folder.",
                  style="CardMuted.TLabel").pack(anchor="w", pady=(4, 0))

        button_row = ttk.Frame(outer)
        button_row.pack(fill="x", pady=(18, 10))
        ttk.Button(button_row, text="Export TXT + Excel", style="Primary.TButton",
                   command=self.export).pack(side="left")
        ttk.Button(button_row, text="Clear", command=self.clear).pack(side="left", padx=(10, 0))

        ttk.Label(outer, textvariable=self.status, style="Sub.TLabel", wraplength=690).pack(anchor="w", pady=(4, 8))

        warning = tk.Frame(outer, bg="#FFF7ED", highlightbackground="#FED7AA", highlightthickness=1)
        warning.pack(fill="x", side="bottom", pady=(8, 0))
        tk.Label(
            warning,
            text="Security note: exports contain readable passwords. Keep them private and delete them securely when finished.",
            bg="#FFF7ED", fg="#9A3412", font=("Segoe UI", 9), wraplength=680,
            justify="left", padx=12, pady=10
        ).pack(fill="x")

    def choose_csv(self):
        path = filedialog.askopenfilename(
            title="Select Chrome Passwords.csv",
            filetypes=[("CSV files", "*.csv"), ("All files", "*.*")]
        )
        if not path:
            return
        self.csv_path.set(path)
        if not self.output_dir.get():
            self.output_dir.set(str(Path(path).parent))
        try:
            self.rows = read_chrome_csv(path)
            self.count_text.set(f"{len(self.rows)} login record(s) found")
            self.status.set("CSV loaded. Choose an output folder, then export.")
        except Exception as exc:
            self.rows = None
            self.count_text.set("Could not read this CSV")
            self.status.set("Please select a valid Chrome password CSV.")
            messagebox.showerror("Could not read CSV", str(exc), parent=self.root)

    def choose_output(self):
        folder = filedialog.askdirectory(title="Choose output folder")
        if folder:
            self.output_dir.set(folder)

    def export(self):
        path = self.csv_path.get().strip()
        folder = self.output_dir.get().strip()
        if not path or not os.path.isfile(path):
            messagebox.showwarning("Choose CSV", "Please choose a valid Chrome Passwords.csv file.", parent=self.root)
            return
        if not folder or not os.path.isdir(folder):
            messagebox.showwarning("Choose folder", "Please choose an existing output folder.", parent=self.root)
            return
        try:
            self.rows = read_chrome_csv(path)
            txt_path = os.path.join(folder, "credentials.txt")
            xlsx_path = os.path.join(folder, "credentials.xlsx")
            export_txt(self.rows, txt_path)
            export_xlsx(self.rows, xlsx_path)
            self.count_text.set(f"{len(self.rows)} login record(s) exported")
            self.status.set(f"Export complete. Files saved in: {folder}")
            messagebox.showinfo(
                "Export complete",
                f"Exported {len(self.rows)} login record(s):\n\n{txt_path}\n{xlsx_path}",
                parent=self.root
            )
        except Exception as exc:
            messagebox.showerror("Export failed", str(exc), parent=self.root)
            self.status.set("Export failed. See the error message for details.")

    def clear(self):
        self.csv_path.set("")
        self.output_dir.set("")
        self.rows = None
        self.count_text.set("No file loaded")
        self.status.set("Choose your Chrome Passwords.csv file to begin.")


def main():
    root = tk.Tk()
    PasswordExporterApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
